#!/usr/bin/env python
"""
Отладка конвертера
"""
from pathlib import Path
from docx import Document
import sys


def debug_docx(file_path):
    """Выводит информацию о DOCX файле"""
    try:
        doc = Document(file_path)

        print(f"Файл: {file_path}")
        print(f"Количество параграфов: {len(doc.paragraphs)}")
        print("\nПервые 10 параграфов:")
        print("-" * 50)

        for i, para in enumerate(doc.paragraphs[:10]):
            style = para.style.name if para.style else "No style"
            text = para.text[:100] if para.text else "[empty]"
            print(f"{i + 1}. [{style}] {text}")

        print("\n" + "=" * 50)

        # Проверяем стили
        styles = set()
        for para in doc.paragraphs:
            if para.style:
                styles.add(para.style.name)

        print(f"\nИспользуемые стили: {sorted(styles)}")

        return True

    except Exception as e:
        print(f"Ошибка: {e}")
        import traceback
        traceback.print_exc()
        return False


if __name__ == "__main__":
    if len(sys.argv) > 1:
        file_path = sys.argv[1]
    else:
        file_path = "/Users/ivt/Documents/ТЕХНИЧЕСКОЕ ЗАДАНИЕ_Тихонов.docx"

    debug_docx(Path(file_path))