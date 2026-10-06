"""Smoke-проверки: python -m pytest tests  (или python tests/test_smoke.py)"""
import os
import sys
import threading
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def test_parallel_threads_get_distinct_libreoffice_profiles():
    from src.core.libreoffice_manager import LibreOfficeManager
    lo = LibreOfficeManager()
    seen, gate = [], threading.Barrier(3)

    def grab():
        seen.append(lo._get_user_profile_path())
        gate.wait()  # держим все три потока живыми одновременно

    threads = [threading.Thread(target=grab) for _ in range(3)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert len(set(seen)) == 3, seen
    # завершившиеся потоки освобождают слоты — новый поток берёт существующий
    t = threading.Thread(target=lambda: seen.append(lo._get_user_profile_path()))
    t.start(); t.join()
    assert seen[-1] in seen[:3]


def test_stylesheets_have_no_unfilled_tokens():
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from src.gui.styles import get_stylesheet_for
    for mode in ("light", "dark", "system"):
        assert "$" not in get_stylesheet_for(mode)


def test_unpack_deb_extracts_files_and_skips_absolute_symlinks(tmp_path=None):
    import io
    import tarfile
    import tempfile
    from src.core.libreoffice_manager import unpack_deb
    tmp = Path(tmp_path or tempfile.mkdtemp())

    data = io.BytesIO()
    with tarfile.open(fileobj=data, mode="w:xz") as tar:
        body = b"#!/bin/sh\n"
        info = tarfile.TarInfo("./opt/libreoffice9.9/program/soffice")
        info.size, info.mode = len(body), 0o755
        tar.addfile(info, io.BytesIO(body))
        link = tarfile.TarInfo("./usr/bin/libreoffice9.9")
        link.type, link.linkname = tarfile.SYMTYPE, "/opt/libreoffice9.9/program/soffice"
        tar.addfile(link)

    def member(name, payload):
        head = f"{name:<16}{0:<12}{0:<6}{0:<6}{100644:<8}{len(payload):<10}`\n".encode()
        return head + payload + (b"\n" if len(payload) % 2 else b"")

    deb = tmp / "pkg.deb"
    deb.write_bytes(b"!<arch>\n" + member("debian-binary", b"2.0\n")
                    + member("control.tar.xz", b"x") + member("data.tar.xz", data.getvalue()))
    unpack_deb(deb, tmp / "out")

    soffice = tmp / "out/opt/libreoffice9.9/program/soffice"
    assert soffice.read_bytes() == b"#!/bin/sh\n" and os.access(soffice, os.X_OK)
    assert not (tmp / "out/usr/bin/libreoffice9.9").exists()   # абсолютная ссылка пропущена


def test_missing_linux_libraries_are_recognised():
    from src.core.libreoffice_manager import missing_libraries
    # Настоящая ошибка из чистого Debian-контейнера
    stderr = ("/root/.local/share/FileConverterPro/libreoffice/linux/opt/libreoffice26.2/program/oosplash: "
              "error while loading shared libraries: libXinerama.so.1: cannot open shared object file: "
              "No such file or directory")
    assert missing_libraries(stderr) == ["libXinerama.so.1"]
    assert missing_libraries("Error: source file could not be loaded") == []


def test_libreoffice_archive_checksum():
    import hashlib
    import tempfile
    from src.core import libreoffice_manager as lm
    f = Path(tempfile.mkdtemp()) / "archive.tar.gz"
    f.write_bytes(b"not really libreoffice")
    good = hashlib.sha256(b"not really libreoffice").hexdigest()
    old = dict(lm.LIBREOFFICE_SHA256)
    try:
        lm.LIBREOFFICE_SHA256["archive.tar.gz"] = good
        assert lm.sha256_matches(f, "archive.tar.gz")
        lm.LIBREOFFICE_SHA256["archive.tar.gz"] = "0" * 64           # подменённый архив
        assert not lm.sha256_matches(f, "archive.tar.gz")
        assert not lm.sha256_matches(f, "unknown.tar.gz")            # неизвестному не доверяем
    finally:
        lm.LIBREOFFICE_SHA256.clear()
        lm.LIBREOFFICE_SHA256.update(old)
    # у каждого архива, который приложение может скачать, есть сумма
    for name in ("libreoffice_macos_arm64.tar.gz", "libreoffice_macos.tar.gz", "libreoffice_windows.zip",
                 "libreoffice_linux.tar.gz", "LibreOffice_26.2.2_Linux_aarch64_deb.tar.gz"):
        assert len(lm.LIBREOFFICE_SHA256[name]) == 64, name


def test_file_size_is_russian():
    from src.utils.helpers import format_size
    assert [format_size(n) for n in (183, 4300, 12_500_000)] == ["183 Б", "4,2 КБ", "11,9 МБ"]


def test_slow_libreoffice_start_is_not_unavailable():
    """Первый запуск свежескачанного LibreOffice бывает дольше тайм-аута проверки (антивирус сканирует
    файлы). Это не поломка: иначе Word → PDF не работает до перезапуска программы."""
    import subprocess
    from unittest import mock
    from src.core.libreoffice_manager import LibreOfficeManager
    lo = LibreOfficeManager()
    lo._soffice_path = Path(sys.executable)                 # любой существующий файл
    with mock.patch("subprocess.run", side_effect=subprocess.TimeoutExpired("soffice", 10)), \
         mock.patch("threading.Thread", lambda target, **kw: mock.Mock(start=target)):
        lo._check_availability()
    assert lo.is_available()


def test_interrupted_libreoffice_install_is_not_used():
    """Программу закрыли посреди распаковки LibreOffice: недоделанный не считается установленным,
    а перед новой установкой его остатки убираются (профиль — оставляем)."""
    import platform
    import tempfile
    from unittest import mock
    from src.core.libreoffice_manager import INSTALLING, LibreOfficeManager
    root = Path(tempfile.mkdtemp())
    soffice = {"Darwin": root / "LibreOffice.app" / "Contents" / "MacOS" / "soffice",
               "Windows": root / "windows" / "LibreOffice" / "program" / "soffice.exe"}.get(
        platform.system(), root / "linux" / "opt" / "libreoffice26.2" / "program" / "soffice")
    soffice.parent.mkdir(parents=True)
    soffice.write_text("#!/bin/sh\n")
    soffice.chmod(0o755)                                        # _find_soffice берёт только исполняемый
    (root / "profile").mkdir()
    with mock.patch.object(LibreOfficeManager, "get_app_support_dir", staticmethod(lambda: root)):
        lo = LibreOfficeManager()
        assert lo._find_soffice() == soffice                    # распаковано до конца — берём
        (root / INSTALLING).write_text("x")
        assert lo._find_soffice() != soffice                    # распаковку прервали — не берём
        LibreOfficeManager._remove_partial_install(root)
    assert [p.name for p in root.iterdir()] == ["profile"]


def test_dropped_folder_skips_service_files():
    """Перетащили папку: замки Word («~$…») и «тени» macOS на флешках («._…») — не файлы для конвертации."""
    import tempfile
    from src.gui.main_window import MainWindow
    folder = Path(tempfile.mkdtemp())
    for name in ("фото.jpg", "~$отчёт.docx", "._фото.jpg", ".скрытый.png", "СКАН.PNG", "заметки.md", "скан.tif"):
        (folder / name).write_bytes(b"x")
    # .tif — то же, что .tiff: так сохраняют сканеры
    assert [p.name for p in MainWindow._files_in(folder)] == sorted(["фото.jpg", "СКАН.PNG", "скан.tif"])


def test_read_only_folder_is_detected():
    """Папка без прав на запись: одна понятная ошибка до начала, а не «Permission denied» на каждый файл."""
    import os
    import stat
    import sys
    import tempfile
    from src.utils.helpers import folder_writable
    folder = Path(tempfile.mkdtemp())
    assert folder_writable(folder)
    if sys.platform != "win32":                         # на Windows chmod не снимает права на папку
        os.chmod(folder, stat.S_IRUSR | stat.S_IXUSR)
        try:
            assert not folder_writable(folder)
        finally:
            os.chmod(folder, stat.S_IRWXU)


def test_second_app_copy_does_not_wipe_libreoffice_install():
    """Программу запустили дважды: пока первая копия ставит LibreOffice (держит install.lock),
    вторая не трогает папку — иначе стёрла бы распаковку первой как «прерванную»."""
    import tempfile
    from unittest import mock
    from PySide6.QtCore import QLockFile
    from src.core.libreoffice_manager import INSTALLING, LibreOfficeManager
    root = Path(tempfile.mkdtemp())
    (root / INSTALLING).write_text("x")                     # первая копия распаковывает
    (root / "LibreOffice.app").mkdir()
    first = QLockFile(str(root / "install.lock"))
    assert first.tryLock(0)
    lo = LibreOfficeManager()       # создаём до подмен: Shiboken падает, если класс с сигналами менять раньше
    try:
        with mock.patch.object(LibreOfficeManager, "get_app_support_dir", staticmethod(lambda: root)), \
             mock.patch.object(lo, "_libreoffice_archive_url") as url:
            lo._do_install_blocking()
        assert not url.called                               # не начала скачивать
        assert (root / "LibreOffice.app").exists()          # и ничего не стёрла
    finally:
        first.unlock()


def test_external_programs_get_system_environment():
    """Собранное приложение подставляет себе свои библиотеки и плагины Qt. Файловый менеджер (Dolphin —
    тоже на Qt) и LibreOffice с ними падают: им — окружение, как до запуска приложения."""
    from unittest import mock
    from src.utils import helpers
    env = {"LD_LIBRARY_PATH": "/app/_internal", "LD_LIBRARY_PATH_ORIG": "/opt/lib", "HOME": "/home/u",
           "QT_PLUGIN_PATH": "/app/_internal/PySide6/Qt/plugins", "QML2_IMPORT_PATH": "/app/_internal/qml"}
    with mock.patch.dict(os.environ, env, clear=True), mock.patch.object(sys, "frozen", True, create=True):
        assert helpers.child_env() == {"LD_LIBRARY_PATH": "/opt/lib", "HOME": "/home/u"}
        del os.environ["LD_LIBRARY_PATH_ORIG"]                  # до запуска переменной не было
        assert helpers.child_env() == {"HOME": "/home/u"}
        with mock.patch("subprocess.Popen") as popen, mock.patch.object(sys, "platform", "linux"):
            helpers.open_folder("/home/u/Загрузки")
        assert popen.call_args.args[0] == ["xdg-open", "/home/u/Загрузки"]
        assert popen.call_args.kwargs["env"] == {"HOME": "/home/u"}


def test_failed_conversion_leaves_no_partial_file():
    """Место на флешке кончилось посреди записи: обрезанный файл с «готовым» именем не оставляем."""
    import tempfile
    from unittest import mock
    from src.core.job_manager import JobManager, JobStatus
    d = Path(tempfile.mkdtemp())
    (d / "отчёт.png").write_bytes(b"x")

    class Broken:
        def convert(self, src, out):
            out.write_bytes(b"\x89PNG\r\n")              # начало файла — и всё
            raise OSError(28, "No space left on device")

    jm = JobManager()
    job = jm.get_job(jm.add_job(d / "отчёт.png", d / "отчёт.jpg", "png", "jpg"))
    with mock.patch("src.core.job_manager.ConverterFactory.get_converter", return_value=Broken()):
        jm._process_job(job)
    assert job.get_status() == JobStatus.FAILED
    assert not (d / "отчёт.jpg").exists()


def test_busy_libreoffice_profile_falls_back_to_own_profile():
    """Профиль LibreOffice занят другим soffice (вторая копия программы, оставшийся после закрытия или тайм-аута):
    soffice молча отдаёт задачу ему и выходит с кодом 0 — а файл так и не появляется. Повторяем на своём профиле."""
    import subprocess
    import tempfile
    from unittest import mock
    from src.core.libreoffice_manager import LibreOfficeManager
    lo = LibreOfficeManager()
    d = Path(tempfile.mkdtemp())
    (d / "отчёт.docx").write_bytes(b"x")
    profiles = []

    def soffice(cmd, **kwargs):
        profiles.append(cmd[1])
        if len(profiles) == 1:                          # занятый профиль: ни вывода, ни файла
            return subprocess.CompletedProcess(cmd, 0, "", "")
        (Path(cmd[cmd.index("--outdir") + 1]) / "отчёт.pdf").write_bytes(b"%PDF-")
        return subprocess.CompletedProcess(cmd, 0, "convert ... using filter : writer_pdf_Export\n", "")

    with mock.patch.object(lo, "_check_attempted", True), mock.patch.object(lo, "_is_available", True), \
         mock.patch.object(lo, "_soffice_path", Path(sys.executable)), mock.patch("subprocess.run", soffice):
        assert lo.convert(d / "отчёт.docx", d / "отчёт.pdf")
    assert (d / "отчёт.pdf").read_bytes() == b"%PDF-"
    assert len(profiles) == 2 and profiles[0] != profiles[1]


def test_csv_opens_in_local_excel():
    """Русский Excel открывает CSV по «;», а «1.5» читает как 1 мая: пишем CSV так, как ждёт Excel системы."""
    import tempfile
    from datetime import datetime
    from unittest import mock
    from openpyxl import Workbook
    from PySide6.QtCore import QLocale
    from src.converters import spreadsheet_converter as sc
    from src.converters.pdf_to_spreadsheet import PdfToSpreadsheetConverter
    assert QLocale(QLocale.Language.Russian, QLocale.Country.Russia).decimalPoint() == ","
    d = Path(tempfile.mkdtemp())
    wb = Workbook()
    wb.active.append(["товар", "цена", "шт", "дата"])
    wb.active.append(["яблоко", 1.5, 3, datetime(2026, 10, 7)])         # дата — без «00:00:00»
    wb.save(d / "t.xlsx")
    for delimiter, line in ((";", "яблоко;1,5;3;2026-10-07"), (",", "яблоко,1.5,3,2026-10-07")):
        with mock.patch.object(sc, "excel_csv_delimiter", return_value=delimiter):
            assert sc.SpreadsheetConverter().convert(d / "t.xlsx", d / f"t{delimiter}.csv")
            assert PdfToSpreadsheetConverter()._save_rows([["а", "б"]], d / f"p{delimiter}.csv", "csv")
        assert (d / f"t{delimiter}.csv").read_text(encoding="utf-8-sig").splitlines()[1] == line
        assert (d / f"p{delimiter}.csv").read_text(encoding="utf-8-sig").strip() == f"а{delimiter}б"


def test_closing_during_libreoffice_download_does_not_crash():
    """Первый запуск: LibreOffice качается несколько минут, программу закрывают — Qt обрывал процесс
    аварийно («QThread: Destroyed while thread is still running»)."""
    import subprocess
    code = f"""
import os, sys, time, tempfile
from pathlib import Path
sys.path.insert(0, {str(Path(__file__).resolve().parent.parent)!r})
from PySide6.QtCore import QCoreApplication, QTimer
app = QCoreApplication([])
from src.core.libreoffice_manager import LibreOfficeManager
lo = LibreOfficeManager()
root = Path(tempfile.mkdtemp())
LibreOfficeManager.get_app_support_dir = staticmethod(lambda: root)
lo._is_available = False
lo._libreoffice_archive_url = lambda: "https://example.invalid/libreoffice.tar.gz"
lo._download_with_progress = lambda url, dest: time.sleep(30)      # «скачивается»
lo.start_auto_install()
QTimer.singleShot(500, app.quit)                                    # пользователь закрыл программу
sys.exit(app.exec())
"""
    r = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, timeout=25)
    assert r.returncode == 0, (r.returncode, r.stderr[-500:])


def test_libreoffice_failure_is_explained():
    """LibreOffice не смог открыть файл: пользователю — понятная причина, а не «Конвертация не удалась»,
    и без лишнего перезапуска на временном профиле (это не «профиль занят»)."""
    import subprocess
    import tempfile
    import zipfile
    from unittest import mock
    from src.core.job_manager import JobManager
    from src.core.libreoffice_manager import LibreOfficeManager
    lo = LibreOfficeManager()
    d = Path(tempfile.mkdtemp())
    with zipfile.ZipFile(d / "таблица.xlsx", "w") as z:          # zip, но внутри не таблица
        z.writestr("[Content_Types].xml", "<broken")
    calls = []

    def soffice(cmd, **kwargs):
        calls.append(cmd)
        return subprocess.CompletedProcess(cmd, 0, "", "Error: source file could not be loaded\n")

    jm = JobManager()
    job = jm.get_job(jm.add_job(d / "таблица.xlsx", d / "таблица.pdf", "xlsx", "pdf"))
    with mock.patch.object(lo, "_check_attempted", True), mock.patch.object(lo, "_is_available", True), \
         mock.patch.object(lo, "_soffice_path", Path(sys.executable)), mock.patch("subprocess.run", soffice):
        jm._process_job(job)
    assert "не смог открыть таблица.xlsx" in job.error_message, job.error_message
    assert len(calls) == 1


def test_pdf_text_is_escaped_in_html():
    """«a < b» из PDF ломал HTML-страницу, а <script> из PDF исполнялся в браузере."""
    import tempfile
    import fitz
    from src.converters.pdf_to_html import PdfToHtmlConverter
    d = Path(tempfile.mkdtemp())
    doc = fitz.open()
    doc.new_page().insert_text((50, 100), 'a < b & <script>alert(1)</script>', fontname="helv")
    doc.save(d / "x.pdf")
    assert PdfToHtmlConverter().convert(d / "x.pdf", d / "x.html")
    page = (d / "x.html").read_text(encoding="utf-8")
    assert "<script>" not in page and "a &lt; b &amp; &lt;script&gt;" in page


def test_utf16_text_and_csv_are_read():
    """«Текст Юникод» из Excel и вывод PowerShell «> файл.txt» — UTF-16: были кашей и «CSV повреждён»."""
    import tempfile
    from openpyxl import load_workbook
    from src.converters.spreadsheet_converter import SpreadsheetConverter
    from src.utils.helpers import read_text_any
    d = Path(tempfile.mkdtemp())
    (d / "ps.txt").write_bytes("Привет, конвертер!\r\n".encode("utf-16"))
    assert read_text_any(d / "ps.txt") == "Привет, конвертер!\r\n"
    (d / "excel.csv").write_bytes("Товар\tКол-во\r\nяблоко\t10\r\n".encode("utf-16"))
    assert SpreadsheetConverter().convert(d / "excel.csv", d / "excel.xlsx")
    assert list(load_workbook(d / "excel.xlsx").active.iter_rows(values_only=True))[1] == ("яблоко", 10)


if __name__ == "__main__":
    test_utf16_text_and_csv_are_read()
    test_pdf_text_is_escaped_in_html()
    test_libreoffice_failure_is_explained()
    test_closing_during_libreoffice_download_does_not_crash()
    test_csv_opens_in_local_excel()
    test_busy_libreoffice_profile_falls_back_to_own_profile()
    test_failed_conversion_leaves_no_partial_file()
    test_external_programs_get_system_environment()
    test_second_app_copy_does_not_wipe_libreoffice_install()
    test_read_only_folder_is_detected()
    test_dropped_folder_skips_service_files()
    test_interrupted_libreoffice_install_is_not_used()
    test_slow_libreoffice_start_is_not_unavailable()
    test_libreoffice_archive_checksum()
    test_file_size_is_russian()
    test_missing_linux_libraries_are_recognised()
    test_unpack_deb_extracts_files_and_skips_absolute_symlinks()
    test_parallel_threads_get_distinct_libreoffice_profiles()
    test_stylesheets_have_no_unfilled_tokens()
    print("ok")
