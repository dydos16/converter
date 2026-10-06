"""
Сквозная проверка на настоящей ОС. В CI запускается на Windows, работает и на macOS/Linux:
окно приложения без экрана (offscreen), вкладки и меню, конвертация PNG → JPG через интерфейс,
автоустановка LibreOffice с нуля (если его ещё нет) и DOCX → PDF.

    python tests/integration_check.py [папка_для_скриншотов]

Код возврата 0 — всё прошло, 1 — что-то сломано.
"""
import os
import sys
import tempfile
import time
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from PySide6.QtWidgets import QApplication, QFileDialog  # noqa: E402

results: list[tuple[bool, str]] = []


def check(name: str, ok: bool, detail: str = ""):
    results.append((ok, name))
    print(f"[{'OK' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""), flush=True)


def wait(app, seconds: float, until=None):
    end = time.time() + seconds
    while time.time() < end:
        app.processEvents()
        if until and until():
            return True
        time.sleep(.01)
    return bool(until and until())


def main() -> int:
    shots = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    if shots:
        shots.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix="fcp_check_"))

    app = QApplication([])
    app.setStyle("Fusion")
    from src.core.settings import Settings
    Settings.save = lambda self: None              # не трогаем настройки пользователя
    from src.gui.main_window import MainWindow
    from src.gui.styles import get_palette_for, get_stylesheet_for
    from src.gui.glass import PROGRESS_ROLE
    from src.converters.factory import ConverterFactory

    w = MainWindow()
    finished = []                                  # итоги установки LibreOffice
    w.libreoffice_manager.install_finished.connect(lambda ok, msg: finished.append((ok, msg)))
    w.settings.settings['auto_open_folder'] = False
    w.settings.settings['show_notifications'] = True
    app.setPalette(get_palette_for("light"))
    app.setStyleSheet(get_stylesheet_for("light"))
    w.show()
    wait(app, 1.0)

    try:
        if os.environ.get("FCP_EXPECT_NON_ADMIN") and sys.platform == "win32":
            import ctypes
            check("Запущено от обычного пользователя, без прав администратора",
                  not ctypes.windll.shell32.IsUserAnAdmin(), os.environ.get("USERNAME", ""))

        # 1. Вкладки
        for i in (1, 2, 0):
            w.tabs.setCurrentIndex(i)
            wait(app, .4)
        check("Вкладки переключаются", w.pages.currentIndex() == 0 and w.tabs.currentIndex() == 0)

        # 2. Последняя пара форматов восстанавливается целиком
        w.settings.settings['recent_formats'] = ["pdf→txt"]
        w.load_settings()
        pair = f"{w.input_format_combo.currentText()}→{w.output_format_combo.currentText()}"
        check("Последние форматы восстанавливаются: PDF → TXT", pair == "pdf→txt", pair)

        # Выпадающее меню формата
        w.input_format_combo.setCurrentText("png")
        w.output_format_combo.showPopup()
        wait(app, .5)
        popup = w.output_format_combo._popup
        check("Меню формата открывается", popup is not None and popup.isVisible())
        if popup is not None:
            popup.close()

        # 3. PNG → JPG через интерфейс
        from PIL import Image
        src = []
        for i in range(3):
            path = tmp / f"img_{i}.png"
            Image.new("RGB", (64, 64), (i * 80, 120, 200)).save(path)
            src.append(path)
        other = tmp / "другая папка"                    # тот же «img_0.png» из другой папки
        other.mkdir()
        Image.new("RGB", (64, 64), (250, 20, 20)).save(other / "img_0.png")
        src.append(other / "img_0.png")
        out = tmp / "out"
        out.mkdir()
        QFileDialog.getExistingDirectory = staticmethod(lambda *a, **k: str(out))
        w.settings.settings['image_max_width'] = 32     # «Макс. ширина» в настройках
        w.input_format_combo.setCurrentText("pdf")      # формат, оставшийся с прошлого сеанса
        w.add_paths(src)
        w.add_paths(src[:2])                            # те же файлы перетащили ещё раз — дублей быть не должно
        w.output_format_combo.setCurrentText("jpg")
        w.start_conversion()
        w.input_format_combo.setCurrentText("Все")     # формат переключили посреди конвертации
        check("Кнопка «Остановить» доступна всю конвертацию", w.convert_btn.isEnabled())
        done = wait(app, 60, lambda: all(w.file_list.item(i).data(PROGRESS_ROLE) in (100, -1)
                                         for i in range(w.file_list.count())))
        made = sorted(p.name for p in out.iterdir())
        check("PNG → JPG через интерфейс, одинаковые имена не затирают друг друга",
              done and made == ["img_0 (1).jpg", "img_0.jpg", "img_1.jpg", "img_2.jpg"], str(made))
        widths = {Image.open(out / name).width for name in made}
        check("«Макс. ширина» из настроек применяется", widths == {32}, str(widths))

        # «фото.jpeg» с iPhone и «фото.jpg» — один формат: оба годятся, «Из» — jpg
        wait(app, 5, lambda: not w.conversion_in_progress)
        w.clear_files()
        w.input_format_combo.setCurrentText("Все")
        for name in ("фото.jpeg", "снимок.JPG"):
            Image.new("RGB", (8, 8)).save(tmp / name, "JPEG")
        w.add_paths([tmp / "фото.jpeg", tmp / "снимок.JPG"])
        check(".jpeg и .jpg конвертируются вместе", w.input_format_combo.currentText() == "jpg"
              and w.convert_btn.isEnabled(), w.input_format_combo.currentText())

        if shots:
            wait(app, 1.0)
            w.grab().save(str(shots / "converter_light.png"))
            w.tabs.setCurrentIndex(1)
            wait(app, .6)
            w.grab().save(str(shots / "settings_light.png"))
            app.setPalette(get_palette_for("dark"))
            app.setStyleSheet(get_stylesheet_for("dark"))
            wait(app, .3)
            w.grab().save(str(shots / "settings_dark.png"))
            w.tabs.setCurrentIndex(0)
            wait(app, .6)
            w.grab().save(str(shots / "converter_dark.png"))

        # 4. LibreOffice: автоустановка с нуля и DOCX → PDF
        lo = w.libreoffice_manager
        if lo.is_available():
            check("LibreOffice уже установлен — автоустановку пропускаем", True, str(lo.get_soffice_path()))
        else:
            # Программа запускает скачивание при старте; если оно не удалось (не было интернета),
            # следующая конвертация пробует снова. Здесь при старте не качали — должен начать шаг 3
            started = wait(app, 3, lambda: lo._install_thread is not None)
            check("Конвертация начинает скачивать LibreOffice, если его нет", started)
            t = time.time()
            if not started:
                lo._do_install_blocking()
            wait(app, 900, lambda: finished)
            wait(app, .5)
            ok = bool(finished and finished[-1][0])
            check("Автоустановка LibreOffice", ok, f"{time.time() - t:.0f} с, {finished[-1][1] if finished else 'нет ответа'}")
            lo._check_attempted = False
        check("soffice найден", lo.is_available(), str(lo.get_soffice_path()))

        from docx import Document
        doc = Document()
        doc.add_heading("Integration check", 1)
        doc.add_paragraph("LibreOffice auto-install works on this OS.")
        docx_path = tmp / "check.docx"
        doc.save(docx_path)
        pdf_path = tmp / "check.pdf"
        converter = ConverterFactory.get_converter("docx", "pdf")
        t = time.time()
        ok = converter.convert(docx_path, pdf_path)
        is_pdf = pdf_path.exists() and pdf_path.read_bytes()[:5] == b"%PDF-"
        check("DOCX → PDF через LibreOffice", ok and is_pdf, f"{time.time() - t:.1f} с")
    except Exception:
        traceback.print_exc()
        check("Проверка дошла до конца без исключений", False)
    finally:
        w.job_manager.stop()
        w.worker.quit()
        w.worker.wait(2000)

    failed = [name for ok, name in results if not ok]
    print(f"\nИтог: {len(results) - len(failed)} из {len(results)} проверок прошли" +
          (f"; не прошли: {', '.join(failed)}" if failed else ""))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
