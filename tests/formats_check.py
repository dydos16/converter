"""
Проверка всех поддерживаемых конвертаций (src/config/formats.py) на реалистичных файлах: русский текст,
несколько страниц, заголовок, жирный и курсив, таблица, картинка. Каждая пара конвертируется, результат
проверяется по содержимому — внутри действительно нужный формат (а не, скажем, DOCX с расширением .odt),
и текст доехал. Потом — трудные случаи из жизни: поворот фото из EXIF, прозрачность, CMYK, CSV и TXT
из русского Windows, скан без текста, большой PDF, оформление при PDF → DOCX.

    python tests/formats_check.py

Нужен LibreOffice — тот же, что у приложения (его ставит integration_check.py).
Код возврата 0 — всё прошло, 1 — что-то сломано.
"""
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
import time
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from loguru import logger  # noqa: E402

logger.remove()                                   # логи конвертеров не мешают отчёту

TEXT = "Привет, конвертер! Съешь же ещё этих мягких французских булок — 2026"
TABLE = [["Товар", "Кол-во", "Цена"], ["яблоко", "10", "99,50"], ["ёлка", "20", "120"]]       # «ё» — на ней ошибалась pdfplumber (Calibri, Windows)
FILLER = ("В этом разделе рассказывается о продажах за квартал: объёмы выросли, а склад опустел. " * 6).strip()
IMAGE_FORMATS = ("jpeg", "png", "webp", "bmp", "gif", "tiff")

# Что должно оказаться внутри файла с таким расширением
EXPECTED = {
    "jpg": "jpeg", "jpeg": "jpeg", "png": "png", "webp": "webp", "bmp": "bmp", "gif": "gif", "tiff": "tiff",
    "pdf": "pdf", "docx": "docx", "xlsx": "xlsx", "pptx": "pptx", "odt": "odt", "odp": "odp", "rtf": "rtf",
    "doc": "ole", "xls": "ole", "ppt": "ole",     # старые форматы Office — контейнер OLE2
    "txt": "text", "md": "text", "csv": "text", "html": "html", "json": "json", "xml": "xml",
}
ODF = {"application/vnd.oasis.opendocument.text": "odt",
       "application/vnd.oasis.opendocument.presentation": "odp",
       "application/vnd.oasis.opendocument.spreadsheet": "ods"}


def sniff(p: Path) -> str:
    """Формат по содержимому, а не по расширению."""
    data = p.read_bytes()
    head = data[:16]
    for magic, kind in ((b"%PDF", "pdf"), (b"\xff\xd8\xff", "jpeg"), (b"\x89PNG", "png"), (b"BM", "bmp"),
                        (b"GIF8", "gif"), (b"II*\x00", "tiff"), (b"MM\x00*", "tiff"), (b"{\\rtf", "rtf"),
                        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", "ole")):
        if head.startswith(magic):
            return kind
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return "webp"
    if head.startswith(b"PK"):
        with zipfile.ZipFile(p) as z:
            names = set(z.namelist())
            if "mimetype" in names:
                return ODF.get(z.read("mimetype").decode().strip(), "zip")
            for marker, kind in (("word/document.xml", "docx"), ("xl/workbook.xml", "xlsx"),
                                 ("ppt/presentation.xml", "pptx")):
                if marker in names:
                    return kind
        return "zip"
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return "binary"
    if "<html" in text[:2000].lower():
        return "html"
    try:
        json.loads(text)
        return "json"
    except ValueError:
        pass
    try:
        ET.fromstring(text)
        return "xml"
    except ET.ParseError:
        return "text"


def text_of(p: Path, kind: str) -> str | None:
    """Текст из результата — где его можно достать без LibreOffice; None — не проверяем."""
    if kind == "pdf":       # pdfplumber на PDF с Calibri (Windows) выдаёт «Отчё(cid:5)т» — читаем как приложение;
        import fitz         # его сборку строк проверяют строгие проверки фраз со знаками
        from src.converters.pdf_text import page_text
        with fitz.open(p) as d:
            return "\n".join(page_text(page) for page in d)
    if kind == "docx":
        from docx import Document
        d = Document(p)
        cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
        return "\n".join([par.text for par in d.paragraphs] + cells)
    if kind == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True)
        return " ".join(str(v) for ws in wb for row in ws.iter_rows(values_only=True) for v in row if v is not None)
    if kind == "pptx":
        from pptx import Presentation
        return " ".join(s.text_frame.text for sl in Presentation(p).slides for s in sl.shapes if s.has_text_frame)
    if kind == "rtf":                              # кириллица в RTF — экранированные \uNNNN; управляющие слова убираем
        raw = p.read_text(encoding="latin-1")
        raw = re.sub(r"\\u(-?\d+)(?:\\'[0-9a-fA-F]{2}|\?)?", lambda m: chr(int(m.group(1)) % 65536), raw)
        raw = re.sub(r"\\'[0-9a-fA-F]{2}|\\[a-zA-Z]+-?\d* ?", "", raw)
        return raw.replace("{", "").replace("}", "")
    if kind in ("text", "html", "json", "xml"):
        return p.read_text(encoding="utf-8-sig", errors="replace")
    return None


def lo_convert(soffice: Path, profile: Path, src: Path, ext: str, outdir: Path) -> Path | None:
    subprocess.run([str(soffice), "--headless", "--norestore", f"-env:UserInstallation={profile.as_uri()}",
                    "--convert-to", ext, "--outdir", str(outdir), str(src)], capture_output=True, timeout=300)
    made = outdir / f"{src.stem}.{ext}"
    return made if made.exists() else None


# --------------------------------------------------------------------------- #
#  Входные файлы
# --------------------------------------------------------------------------- #

def photo(w: int, h: int):
    """Картинка «как фото»: плавные переходы и фигуры (шум сделал бы PNG огромным и медленным)."""
    from PIL import Image, ImageDraw
    ramp = Image.linear_gradient("L").resize((w, h))
    img = Image.merge("RGB", (ramp, ramp.transpose(Image.Transpose.ROTATE_90).resize((w, h)), Image.new("L", (w, h), 160)))
    d = ImageDraw.Draw(img)
    d.ellipse([w // 5, h // 5, w // 2, h // 2], fill=(250, 200, 40))
    d.rectangle([w // 2, h // 2, w - w // 6, h - h // 6], fill=(30, 120, 220))
    return img


def rich_docx(path: Path, pages: int, picture: Path):
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    doc.add_heading("Отчёт о продажах", 0)
    p = doc.add_paragraph(TEXT + " ")
    p.add_run("жирный").bold = True
    p.add_run(" и ")
    p.add_run("курсив").italic = True
    table = doc.add_table(rows=len(TABLE), cols=len(TABLE[0]))
    table.style = "Table Grid"
    for r, row in enumerate(TABLE):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    doc.add_picture(str(picture), width=Inches(3))
    for n in range(2, pages + 1):
        doc.add_page_break()
        doc.add_heading(f"Раздел {n}", 1)
        doc.add_paragraph(FILLER)
    doc.save(path)


def make_inputs(d: Path, soffice: Path, profile: Path) -> dict[str, Path]:
    import csv
    import pillow_heif
    from openpyxl import Workbook
    from pptx import Presentation

    d.mkdir(parents=True)
    lo = d / "lo"
    img = photo(2400, 1600)
    files = {}
    for ext, fmt in (("png", "PNG"), ("jpg", "JPEG"), ("jpeg", "JPEG"), ("webp", "WEBP"), ("bmp", "BMP"),
                     ("gif", "GIF"), ("tiff", "TIFF")):
        files[ext] = d / f"sample.{ext}"
        img.save(files[ext], fmt)
    pillow_heif.register_heif_opener()
    for ext in ("heic", "heif"):
        files[ext] = d / f"sample.{ext}"
        img.save(files[ext], "HEIF")

    files["docx"] = d / "sample.docx"
    rich_docx(files["docx"], 3, files["png"])

    prs = Presentation()
    for title, body in ((TEXT, "яблоко и груша"), ("Итоги", "Продажи выросли на 12 %")):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    files["pptx"] = d / "sample.pptx"
    prs.save(files["pptx"])

    wb = Workbook()
    wb.active.title = "Продажи"
    wb.active.append(TABLE[0])
    for name, qty, price in TABLE[1:]:
        wb.active.append([name, int(qty), float(price.replace(",", "."))])
    files["xlsx"] = d / "sample.xlsx"
    wb.save(files["xlsx"])

    files["csv"] = d / "sample.csv"
    with open(files["csv"], "w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(TABLE)
    files["txt"] = d / "sample.txt"
    files["txt"].write_text(f"{TEXT}\n{FILLER}\n", encoding="utf-8")

    # PDF — из документа Word через LibreOffice, как у людей; старые и ODF-форматы — тоже им
    for ext, src in (("pdf", "docx"), ("doc", "docx"), ("odt", "docx"), ("rtf", "docx"), ("ppt", "pptx"),
                     ("pps", "pptx"), ("ppsx", "pptx"), ("xls", "xlsx")):
        made = lo_convert(soffice, profile, files[src], ext, lo)
        if made:
            files[ext] = made
    return files


# --------------------------------------------------------------------------- #
#  Проверки
# --------------------------------------------------------------------------- #

def run(src: Path, out: Path, **kwargs) -> tuple[bool, list[str]]:
    from src.converters.factory import ConverterFactory
    errors: list[str] = []
    converter = ConverterFactory.get_converter(src.suffix, out.suffix, **kwargs)
    converter.error_callback = errors.append
    out.parent.mkdir(parents=True, exist_ok=True)
    return bool(converter.convert(src, out)), errors


def why(errors: list[str]) -> str:
    return f"ошибка: {errors[-1][:200]}" if errors else "конвертер вернул ошибку"


def check_pair(src: str, dst: str, files: dict[str, Path], out: Path, pages: int) -> str | None:
    """None — всё хорошо, иначе — что не так."""
    ok, errors = run(files[src], out)
    if not ok:
        return why(errors)
    want = EXPECTED[dst]
    if not out.exists() and src == "pdf" and want in IMAGE_FORMATS and pages > 1:
        # Многостраничный PDF → картинки: архив со всеми страницами
        archive = out.with_suffix(".zip")
        if not archive.exists():
            return "нет архива со страницами"
        with zipfile.ZipFile(archive) as z:
            names = z.namelist()
            if len(names) != pages:
                return f"в архиве {len(names)} страниц из {pages}"
            out = out.parent / names[0]
            out.write_bytes(z.read(names[0]))
    if not out.exists():
        others = [p.name for p in out.parent.iterdir()]
        return "нет файла результата" + (f" (есть: {', '.join(others)})" if others else "")
    kind = sniff(out)
    if kind != want:
        return f"внутри {kind.upper()}, а не {want.upper()}"
    # Текст должен доехать везде, кроме картинок (в PDF → PPTX он теперь настоящий, редактируемый)
    if src in ("png", "jpg", "jpeg", "webp", "bmp", "gif", "tiff", "heic", "heif") or want in IMAGE_FORMATS:
        return None
    # Фраза целиком, со знаками: так ловится и испорченная кириллица, и разъехавшиеся запятые
    markers = ("яблоко", "Кол-во", "ёлка") if {src, dst} & {"xlsx", "xls", "csv"} else ("Привет, конвертер!",)
    text = text_of(out, kind)
    if text is not None:
        flat = re.sub(r"\s+", " ", text)
        missing = [m for m in markers if m not in flat]
        if missing:
            i = flat.find(missing[0][:4])
            return f"в результате нет «{missing[0]}»" + (f": …{flat[max(0, i - 30):i + 50]}…" if i >= 0 else "")
    return None


def edge_cases(files: dict[str, Path], d: Path, soffice: Path, profile: Path) -> list:
    """Трудные случаи из жизни: (название, функция → None или описание проблемы)."""
    from PIL import Image
    d.mkdir(parents=True)

    def exif_rotation():
        src = d / "rotated.jpg"
        exif = Image.Exif()
        exif[0x0112] = 6                                # «повернуть на 90°» — так снимает телефон
        photo(300, 200).save(src, "JPEG", exif=exif)
        ok, err = run(src, d / "rotated.png")
        size = Image.open(d / "rotated.png").size if ok else None
        return None if size == (200, 300) else (why(err) if not ok else f"размер {size} — фото лежит на боку")

    def transparent_to_jpg(palette: bool):
        def check():
            src = d / f"transparent_{palette}.png"
            im = Image.new("RGBA", (200, 100), (0, 0, 0, 0))
            im.paste((220, 30, 30, 255), (100, 0, 200, 100))
            if palette:
                im = im.convert("P", palette=Image.Palette.ADAPTIVE, colors=8)
                im.info["transparency"] = im.getpixel((10, 10))
                im.save(src, transparency=im.info["transparency"])
            else:
                im.save(src)
            out = d / f"transparent_{palette}.jpg"
            ok, err = run(src, out)
            if not ok:
                return why(err)
            pixel = Image.open(out).convert("RGB").getpixel((20, 50))
            return None if min(pixel) > 230 else f"прозрачный фон стал {pixel}, а не белым"
        return check

    def cmyk(dst: str):
        def check():
            src = d / "cmyk.jpg"
            photo(300, 200).convert("CMYK").save(src, "JPEG")
            ok, err = run(src, d / f"cmyk.{dst}")
            return None if ok and sniff(d / f"cmyk.{dst}") == EXPECTED[dst] else why(err)
        return check

    def png16():
        src = d / "deep.png"
        Image.linear_gradient("L").resize((300, 200)).convert("I;16").save(src)
        ok, err = run(src, d / "deep.jpg")
        return None if ok else why(err)

    def windows_csv(dst: str):
        def check():
            src = d / f"excel_ru_{dst}.csv"             # «CSV (разделители — точка с запятой)» из русского Excel
            src.write_bytes("\r\n".join(";".join(r) for r in TABLE).encode("cp1251") + b"\r\n")
            out = d / f"excel_ru.{dst}"
            ok, err = run(src, out)
            if not ok:
                return why(err)
            if dst == "xlsx":
                from openpyxl import load_workbook
                ws = load_workbook(out).active
                return None if ws["A2"].value == "яблоко" and ws["B2"].value in (10, "10") else \
                    f"ячейки A2={ws['A2'].value!r} B2={ws['B2'].value!r} — кодировка или разделитель не распознаны"
            return None if "яблоко" in (text_of(out, sniff(out)) or "") else "в PDF нет «яблоко»"
        return check

    def windows_txt(dst: str):
        def check():
            src = d / f"notepad_{dst}.txt"
            src.write_bytes(f"{TEXT}\r\n".encode("cp1251"))
            out = d / f"notepad.{dst}"
            ok, err = run(src, out)
            if not ok:
                return why(err)
            return None if "Привет" in (text_of(out, sniff(out)) or "") else "кириллица испорчена"
        return check

    def csv_for_excel():
        out = d / "for_excel.csv"
        ok, err = run(files["xlsx"], out)
        if not ok:
            return why(err)
        return None if out.read_bytes().startswith(b"\xef\xbb\xbf") else "CSV без BOM — русский Excel покажет кракозябры"

    def scan_pdf() -> Path:
        """Скан как со сканера: страница картинкой 200 DPI, текстового слоя нет."""
        import fitz
        src = d / "scan.pdf"
        if not src.exists():
            with fitz.open(files["pdf"]) as pdf:
                png = pdf[0].get_pixmap(dpi=200).tobytes("png")
            Image.open(io.BytesIO(png)).convert("RGB").save(src, "PDF", resolution=200)
        return src

    def scan_to_text():
        out = d / "scan.txt"
        ok, err = run(scan_pdf(), out)
        if not ok:
            return why(err)
        flat = re.sub(r"\s+", " ", out.read_text(encoding="utf-8-sig"))
        missing = [m for m in ("Привет, конвертер!", "яблоко", "Кол-во") if m not in flat]
        return None if not missing else f"не распознано: {', '.join(missing)} (вышло: {flat[:120]}…)"

    def scan_to_docx():
        out = d / "scan.docx"
        ok, err = run(scan_pdf(), out)
        if not ok:
            return why(err)
        from docx import Document
        flat = re.sub(r"\s+", " ", " ".join(p.text for p in Document(out).paragraphs))
        return None if "Привет, конвертер!" in flat else f"текст не распознан: {flat[:120]}…"

    def pdf_to_pptx_editable():
        out = d / "slides.pptx"
        ok, err = run(files["pdf"], out)
        if not ok:
            return why(err)
        from pptx import Presentation
        slide = Presentation(out).slides[0]
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
        pictures = [s for s in slide.shapes if s.shape_type == 13]          # 13 — картинка (фон страницы)
        missing = [what for what, present in (
            ("заголовок надписью", "Отчёт о продажах" in texts),
            ("абзац надписью", any(TEXT in re.sub(r"\s+", " ", t) for t in texts)),
            ("ячейки таблицы на своих местах", "Кол-во" in texts and "яблоко" in texts),
            ("фон страницы", len(pictures) == 1),
        ) if not present]
        return None if not missing else "нет: " + ", ".join(missing)

    def pdf_to_docx_keeps_formatting():
        out = d / "formatted.docx"
        ok, err = run(files["pdf"], out)
        if not ok:
            return why(err)
        from docx import Document
        doc = Document(out)
        runs = [r for p in doc.paragraphs for r in p.runs]
        missing = [what for what, present in (
            ("таблица", any("яблоко" in c.text for t in doc.tables for row in t.rows for c in row.cells)),
            ("«Кол-во» в шапке таблицы", any(c.text == "Кол-во" for t in doc.tables for row in t.rows for c in row.cells)),
            ("текст абзаца по порядку", TEXT in re.sub(r"\s+", " ", " ".join(p.text for p in doc.paragraphs))),
            ("жирный", any(r.bold and "жирный" in r.text for r in runs)),
            ("курсив", any(r.italic and "курсив" in r.text for r in runs)),
            ("картинка", bool(doc.inline_shapes)),
            ("крупный заголовок", any("Отчёт" in r.text and r.font.size and r.font.size.pt >= 18 for r in runs)),
        ) if not present]
        return None if not missing else "потеряно: " + ", ".join(missing)

    def big_pdf_to_png():
        import fitz
        big_docx = d / "big.docx"
        rich_docx(big_docx, 40, files["png"])
        big = lo_convert(soffice, profile, big_docx, "pdf", d)
        with fitz.open(big) as pdf:
            pages = len(pdf)
        try:
            import resource                             # на Windows пиковую память так не померить
            scale = 1 if sys.platform == "darwin" else 1024     # ru_maxrss: байты на Mac, КБ на Linux
            peak = lambda: resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * scale
        except ImportError:
            peak = lambda: 0
        before = peak()
        ok, err = run(big, d / "big.png", dpi=200)
        grown = (peak() - before) / 1024 ** 2
        if not ok:
            return why(err)
        with zipfile.ZipFile(d / "big.zip") as z:
            n = len(z.namelist())
        if n != pages:
            return f"в архиве {n} страниц из {pages}"
        return None if grown < 300 else f"память выросла на {grown:.0f} МБ — страницы копятся в памяти"

    def page_order():
        import fitz
        ok, err = run(files["pdf"], d / "order.png", dpi=72)
        if not ok:
            return why(err)
        with zipfile.ZipFile(d / "order.zip") as z, fitz.open(files["pdf"]) as pdf:
            for n in range(1, len(pdf) + 1):
                got = Image.open(io.BytesIO(z.read(f"order_page_{n}.png"))).convert("RGB").tobytes()
                if got != pdf[n - 1].get_pixmap(dpi=72, alpha=False).samples:
                    return f"в файле страницы {n} лежит другая страница"
        return None

    def password_pdf():
        import fitz
        from src.core.job_manager import ConversionJob, JobManager
        src = d / "secret.pdf"
        with fitz.open(files["pdf"]) as doc:
            doc.save(src, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="1234", owner_pw="1234")
        job = ConversionJob("secret", src, d / "secret.txt", "pdf", "txt")
        JobManager()._process_job(job)                  # как в приложении: через очередь задач
        return None if "паролем" in job.error_message else f"сообщение: {job.error_message!r}"

    def compress_scan():
        import fitz
        src = d / "scan300.pdf"
        with fitz.open(files["pdf"]) as pdf:
            png = pdf[0].get_pixmap(dpi=300).tobytes("png")
        Image.open(io.BytesIO(png)).convert("RGB").save(src, "PDF", resolution=300, quality=95)
        out = d / "scan300_small.pdf"
        ok, err = run(src, out, compression_level=6)
        if not ok:
            return why(err)
        before, after = src.stat().st_size, out.stat().st_size
        return None if after < before * 0.7 else f"{before // 1024} КБ → {after // 1024} КБ — почти не сжалось"

    def two_columns_to_docx():
        from docx import Document
        from docx.enum.text import WD_BREAK
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        src_docx = d / "columns.docx"
        doc = Document()
        cols = OxmlElement("w:cols")
        cols.set(qn("w:num"), "2")
        doc.sections[0]._sectPr.append(cols)
        for i in range(1, 6):
            doc.add_paragraph(f"Колонка один, абзац {i}. {FILLER[:120]}")
        doc.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
        for i in range(1, 6):
            doc.add_paragraph(f"Колонка два, абзац {i}. {FILLER[:120]}")
        doc.save(src_docx)
        pdf = lo_convert(soffice, profile, src_docx, "pdf", d)
        out = d / "columns_back.docx"
        ok, err = run(pdf, out)
        if not ok:
            return why(err)
        order = [("один" if "Колонка один" in p.text else "два") for p in Document(out).paragraphs
                 if "Колонка" in p.text]
        return None if order == ["один"] * 5 + ["два"] * 5 else f"порядок абзацев: {' '.join(order)}"

    def rotated_to_pptx():
        import fitz
        from pptx import Presentation
        src = d / "rotated.pdf"
        with fitz.open(files["pdf"]) as doc:
            doc[0].set_rotation(90)
            doc.save(src)
        out = d / "rotated.pptx"
        ok, err = run(src, out)
        if not ok:
            return why(err)
        boxes = [s for s in Presentation(out).slides[0].shapes if s.has_text_frame]
        return None if not boxes else f"{len(boxes)} надписей легли мимо повёрнутой страницы"

    def existing_file_kept():
        from src.utils.helpers import get_unique_filename
        out_dir = d / "Загрузки"
        out_dir.mkdir()
        mine = out_dir / "sample.pdf"
        mine.write_bytes(b"%PDF-1.4 MY OWN FILE")       # у пользователя уже лежит файл с этим именем
        target = get_unique_filename(out_dir / "sample.pdf")
        ok, err = run(files["docx"], target)
        if not ok:
            return why(err)
        if not mine.exists() or b"MY OWN FILE" not in mine.read_bytes():
            return "файл пользователя с тем же именем затёрт"
        return None if target.exists() and sniff(target) == "pdf" else "результата нет"

    def same_names_in_parallel():
        import threading
        from docx import Document
        from src.utils.helpers import get_unique_filename
        out_dir, taken, jobs = d / "Итог", set(), []
        out_dir.mkdir()
        for dept in ("А", "Б", "В"):
            src = d / f"отдел {dept}" / "отчёт.docx"     # три «отчёт.docx» из разных папок
            src.parent.mkdir()
            doc = Document()
            doc.add_paragraph(f"Отчёт отдела {dept}")
            doc.save(src)
            target = get_unique_filename(out_dir / "отчёт.pdf", taken)
            taken.add(target)
            jobs.append((dept, src, target))
        results, logs = {}, []
        sink = logger.add(lambda m: logs.append(m.record["message"]), level="WARNING")   # что скажет LibreOffice
        threads = [threading.Thread(target=lambda j=job: results.__setitem__(j[0], run(j[1], j[2]))) for job in jobs]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        logger.remove(sink)
        for message in logs:
            if "код 81" in message:
                print(f"   (LibreOffice: {message})", flush=True)
        def inside(target: Path) -> str:
            return re.sub(r"\s+", " ", text_of(target, "pdf") or "")[:50] if target.exists() else "файла нет"

        wrong = [f"{target.name}: ждали отдел {dept}, внутри «{inside(target)}» ({target.stat().st_size if target.exists() else 0} Б)"
                 for dept, _, target in jobs if f"Отчёт отдела {dept}" not in inside(target)]
        if not wrong:
            return None
        return "; ".join(wrong) + f" | итоги задач: {results} | журнал: {' / '.join(logs)[-700:]}"

    def broken_inputs():
        import os
        from src.core.job_manager import ConversionJob, JobManager
        jm, bad, folder = JobManager(), [], d / "битые"
        folder.mkdir()
        for name, data, dst, expect in (
                ("пустой.docx", b"", "pdf", "пустой"), ("пустой.png", b"", "jpg", "пустой"),
                ("мусор.docx", os.urandom(3000), "pdf", "повреждён"), ("мусор.xlsx", os.urandom(3000), "csv", "повреждён"),
                ("мусор.pptx", os.urandom(3000), "pdf", "повреждён"), ("мусор.pdf", os.urandom(3000), "txt", "повреждён"),
                ("мусор.png", os.urandom(3000), "jpg", "не картинка")):
            src = folder / name
            src.write_bytes(data)
            job = ConversionJob(name, src, folder / f"итог {name}.{dst}", src.suffix.lstrip("."), dst)
            jm._process_job(job)                        # как в приложении: через очередь задач
            if job.status.name != "FAILED" or expect not in job.error_message:
                bad.append(f"{name}: {job.status.name} «{job.error_message[:70]}»")
        gone = ConversionJob("gone", folder / "удалённый.docx", folder / "x.pdf", "docx", "pdf")
        jm._process_job(gone)
        if "не найден" not in gone.error_message:
            bad.append(f"удалённый файл: «{gone.error_message}»")
        return None if not bad else "; ".join(bad)

    def terminal_log_txt():
        src = d / "build log.txt"
        src.write_text("\x1b[32mOK\x1b[0m Сборка прошла\n\x00Ошибок нет\x07\n", encoding="utf-8")
        out = d / "build log.docx"
        ok, err = run(src, out)
        if not ok:
            return why(err)
        text = text_of(out, "docx") or ""
        return None if "OK Сборка прошла" in text and "Ошибок нет" in text else f"текст: {text!r}"

    def odd_file_name():
        import unicodedata
        # Знаки, которые LibreOffice понимает как части адреса, и «й» в разложенной форме, как бывает на Mac
        src = d / unicodedata.normalize("NFD", "Отчёт #1 (копия) 100% & «итог»; й.docx")
        shutil.copy(files["docx"], src)
        out = d / "odd name.pdf"
        ok, err = run(src, out)
        if not ok:
            return why(err)
        return None if "Привет, конвертер!" in re.sub(r"\s+", " ", text_of(out, "pdf") or "") else "в PDF нет текста"

    def multipage_tiff():
        pages = [Image.new("RGB", (200, 100), c) for c in ((200, 0, 0), (0, 200, 0), (0, 0, 200))]
        pages[0].save(d / "fax.tiff", save_all=True, append_images=pages[1:])
        ok, err = run(d / "fax.tiff", d / "fax.png")
        if not ok:
            return why(err)
        with zipfile.ZipFile(d / "fax.zip") as z:
            n = len(z.namelist())
        return None if n == 3 else f"страниц в архиве: {n} из 3"

    def animated_gif():
        frames = [Image.new("RGB", (60, 60), (i * 80, 50, 50)) for i in range(3)]
        frames[0].save(d / "anim.gif", save_all=True, append_images=frames[1:], duration=200, loop=0)
        ok, err = run(d / "anim.gif", d / "anim.webp")
        if not ok:
            return why(err)
        n = getattr(Image.open(d / "anim.webp"), "n_frames", 1)
        return None if n == 3 else f"кадров: {n} из 3"

    def panorama():
        src = d / "pano.png"
        Image.new("L", (13500, 13500), 200).save(src)   # 182 Мпикс — больше порога Pillow
        ok, err = run(src, d / "pano.jpg")
        return None if ok else why(err)

    def sheets_to_csv():
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "Январь"
        wb.active.append(["январь", 1])
        wb.create_sheet("Февраль").append(["февраль", 2])
        wb.create_sheet("Пустой")
        wb.save(d / "months.xlsx")
        ok, err = run(d / "months.xlsx", d / "months.csv")
        if not ok:
            return why(err)
        made = sorted(p.name for p in d.glob("months*.csv"))
        return None if made == ["months - Февраль.csv", "months.csv"] else f"файлы: {made}"

    def image_max_size():
        too_big = []
        for src, dst in (("png", "jpg"), ("heic", "png")):
            out = d / f"small_{src}.{dst}"
            ok, err = run(files[src], out, max_width=100, max_height=None)
            if not ok:
                return why(err)
            if Image.open(out).width > 100:
                too_big.append(f"{src} → {dst}: {Image.open(out).width} px")
        return None if not too_big else "не уменьшено: " + ", ".join(too_big)

    def borderless_table():
        from docx import Document
        src_docx = d / "borderless.docx"
        doc = Document()
        table = doc.add_table(rows=len(TABLE), cols=len(TABLE[0]))   # без стиля — без рамок
        for r, row in enumerate(TABLE):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        doc.save(src_docx)
        pdf = lo_convert(soffice, profile, src_docx, "pdf", d)
        out = d / "borderless.xlsx"
        ok, err = run(pdf, out, table_strategy="stream")
        if not ok:
            return why(err)
        from openpyxl import load_workbook
        rows = [[v for v in row if v not in (None, "")] for row in load_workbook(out).active.iter_rows(values_only=True)]
        hit = next((row for row in rows if "яблоко" in row), None)
        return None if hit and len(hit) >= 3 else f"строка с «яблоко» не разбита по столбцам: {hit}"

    def tiff_compressed():
        ok, err = run(files["pdf"], d / "page.tiff")
        if not ok:
            return why(err)
        with zipfile.ZipFile(d / "page.zip") as z:
            size = max(i.file_size for i in z.infolist())
        return None if size < 3 * 1024 ** 2 else f"TIFF {size / 1024 ** 2:.0f} МБ на страницу — без сжатия"

    return [
        ("PDF 40 страниц → PNG: все страницы, без лишней памяти", big_pdf_to_png),
        ("Фото с поворотом в EXIF → PNG", exif_rotation),
        ("PNG с прозрачностью → JPG, фон белый", transparent_to_jpg(False)),
        ("PNG с палитрой и прозрачностью → JPG, фон белый", transparent_to_jpg(True)),
        ("CMYK JPEG → PNG", cmyk("png")),
        ("CMYK JPEG → WEBP", cmyk("webp")),
        ("PNG 16 бит → JPG", png16),
        ("CSV из русского Excel (1251, «;») → XLSX", windows_csv("xlsx")),
        ("CSV из русского Excel (1251, «;») → PDF", windows_csv("pdf")),
        ("TXT из Блокнота (1251) → DOCX", windows_txt("docx")),
        ("TXT из Блокнота (1251) → PDF", windows_txt("pdf")),
        ("XLSX → CSV открывается в русском Excel (BOM)", csv_for_excel),
        ("Скан PDF → TXT: текст распознан (OCR), таблица тоже", scan_to_text),
        ("Скан PDF → DOCX: текст распознан (OCR)", scan_to_docx),
        ("PDF → PPTX: текст редактируемый, таблица по ячейкам", pdf_to_pptx_editable),
        ("PDF → DOCX сохраняет оформление", pdf_to_docx_keeps_formatting),
        ("PDF → PNG: страницы по порядку", page_order),
        ("PDF → TIFF сжат", tiff_compressed),
        ("PDF под паролем: понятное сообщение", password_pdf),
        ("Сжатие PDF со сканом уменьшает файл", compress_scan),
        ("PDF в две колонки → DOCX: сначала левая, потом правая", two_columns_to_docx),
        ("Повёрнутая страница → PPTX: без надписей мимо", rotated_to_pptx),
        ("Файл пользователя с тем же именем не затирается", existing_file_kept),
        ("Три «отчёт.docx» одновременно → три правильных PDF", same_names_in_parallel),
        ("Многостраничный TIFF → PNG: все страницы", multipage_tiff),
        ("Анимированный GIF → WEBP: анимация сохранена", animated_gif),
        ("Панорама 182 Мпикс → JPG", panorama),
        ("XLSX с несколькими листами → CSV на каждый лист", sheets_to_csv),
        ("«Макс. ширина» уменьшает картинку (и HEIC)", image_max_size),
        ("Таблица без рамок → XLSX (режим stream)", borderless_table),
        ("Пустые, испорченные и удалённые файлы: понятная ошибка, а не «успех»", broken_inputs),
        ("TXT с кодами цветов терминала → DOCX", terminal_log_txt),
        ("Имя с #, %, &, «» и «й» как на Mac → PDF через LibreOffice", odd_file_name),
    ]


# --------------------------------------------------------------------------- #

def main() -> int:
    from src.config.formats import SUPPORTED_CONVERSIONS
    from src.core.libreoffice_manager import LibreOfficeManager

    soffice = LibreOfficeManager()._find_soffice()
    if not soffice:
        print("LibreOffice не найден — сначала запустите tests/integration_check.py (он его ставит)")
        return 1

    # Кириллица и пробелы в пути — как у русских пользователей Windows («C:\Users\Иван\Мои документы»)
    tmp = Path(tempfile.mkdtemp(prefix="fcp форматы "))
    pairs = sorted(SUPPORTED_CONVERSIONS)
    failed, started, cases = [], time.time(), []
    try:
        profile = tmp / "lo-profile"
        files = make_inputs(tmp / "in", soffice, profile)
        import fitz
        with fitz.open(files["pdf"]) as pdf:
            pdf_pages = len(pdf)
        print(f"Пар конвертации: {len(pairs)}, страниц в тестовом PDF: {pdf_pages}; LibreOffice: {soffice}\n",
              flush=True)

        cases = edge_cases(files, tmp / "edge", soffice, profile)
        checks = [(f"{s} → {t}", lambda s=s, t=t: check_pair(s, t, files, tmp / "out" / f"{s}_to_{t}" / f"result.{t}",
                                                             pdf_pages) if s in files else f"не удалось создать входной .{s}")
                  for s, t in pairs]
        # Большой PDF — первым: пиковая память процесса ещё не раздута другими конвертациями
        for name, fn in cases[:1] + checks + cases[1:]:
            t = time.time()
            try:
                problem = fn()
            except Exception as e:                      # noqa: BLE001 — нужен отчёт по каждой проверке
                problem = f"исключение {type(e).__name__}: {e}"
            if problem:
                failed.append(name)
                print(f"[FAIL] {name} — {problem}", flush=True)
            else:
                print(f"[OK] {name} ({time.time() - t:.1f} с)", flush=True)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    total = len(pairs) + len(cases)
    print(f"\nИтог: {total - len(failed)} из {total} проверок прошли за {time.time() - started:.0f} с" +
          (f"; не прошли: {', '.join(failed)}" if failed else ""))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
