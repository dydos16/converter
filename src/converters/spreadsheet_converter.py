"""
Таблицы: xlsx, xls, csv — друг в друга и в PDF.
CSV разбираем сами (UTF-8 или Windows-1251, разделитель «,», «;» или табуляция — как сохраняет русский Excel),
CSV для Excel пишем с BOM. Остальное делает LibreOffice.
"""
import csv
import io
import re
import tempfile
from datetime import datetime, time
from pathlib import Path
from .base import BaseConverter
from loguru import logger


EXCEL_MAX_ROWS = 1_048_576          # больше строк в лист Excel не помещается
# Символы, которые нельзя хранить в ячейке (то же, что openpyxl.cell.cell.ILLEGAL_CHARACTERS_RE;
# сам openpyxl здесь не импортируем — модуль грузится при запуске окна)
ILLEGAL_CHARACTERS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def excel_csv_delimiter() -> str:
    """
    Разделитель, которого ждёт Excel этой системы. Где дробную часть пишут через запятую (Россия), Excel
    открывает CSV по «;»: с «,» вся таблица ложится в один столбец, а «1.5» он читает как 1 мая.
    """
    # ponytail: судим по десятичному знаку, а не по «Разделителю элементов списка» Windows — он почти всегда
    # ему парный; читать его через GetLocaleInfo, если кто-то настроит их вразнобой
    from PySide6.QtCore import QLocale
    return ";" if QLocale.system().decimalPoint() == "," else ","


def _csv_value(value, comma: bool):
    """Ячейка для CSV: дата — без «00:00:00», дробь — через запятую, если так пишут в системе."""
    if isinstance(value, datetime) and value.time() == time():
        return value.date()
    if comma and isinstance(value, float):
        return str(value).replace('.', ',')
    return value


def _cell(value: str):
    """Числа из CSV — числами, иначе Excel покажет их текстом. Коды с ведущим нулём и длинные номера не трогаем."""
    # Управляющие символы (цвета терминала в логах, мусор выгрузок) Excel хранить не умеет — вся таблица падала
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if re.fullmatch(r"-?(0|[1-9]\d{0,14})", value):
        return int(value)
    if re.fullmatch(r"-?(0|[1-9]\d{0,14})[.,]\d+", value):
        return float(value.replace(",", "."))
    return value


class SpreadsheetConverter(BaseConverter):
    """Конвертер таблиц через openpyxl и LibreOffice"""

    def get_input_formats(self):
        return ['xlsx', 'xls', 'csv']

    def get_output_formats(self):
        return ['xlsx', 'xls', 'csv', 'pdf']

    def convert(self, input_path: Path, output_path: Path) -> bool:
        from src.core.libreoffice_manager import LibreOfficeManager, LO_MISSING
        src = input_path.suffix.lower().lstrip('.')
        dst = output_path.suffix.lower().lstrip('.')
        try:
            if src == 'csv' and dst == 'xlsx':
                return self._csv_to_xlsx(input_path, output_path)
            if src == 'xlsx' and dst == 'csv':
                return self._xlsx_to_csv(input_path, output_path)

            lo = LibreOfficeManager()
            if not lo.is_available():
                self._handle_error(LO_MISSING)
                return False
            self._update_status(f"Конвертация через LibreOffice в {dst.upper()}...")
            if src == 'csv':        # → pdf/xls: CSV разбираем сами в XLSX, дальше LibreOffice
                return lo.convert_via(lambda xlsx: self._csv_to_xlsx(input_path, xlsx), output_path, '.xlsx')
            if dst == 'csv':        # xls → csv: LibreOffice делает XLSX, CSV для Excel пишем сами
                with tempfile.TemporaryDirectory() as tmp:
                    xlsx = Path(tmp) / f"{input_path.stem}.xlsx"
                    return lo.convert(input_path, xlsx) and self._xlsx_to_csv(xlsx, output_path)
            return lo.convert(input_path, output_path)      # xlsx ↔ xls, таблица → PDF

        except csv.Error:
            self._handle_error(f"Не удалось разобрать CSV — файл повреждён или это не таблица: {input_path.name}.")
            return False
        except Exception as e:
            self._handle_error(f"Ошибка конвертации таблицы: {str(e)}")
            logger.exception("Ошибка конвертации таблицы")
            return False

    def _csv_to_xlsx(self, input_path: Path, output_path: Path) -> bool:
        from openpyxl import Workbook
        from src.utils.helpers import read_text_any

        self._update_status(f"Чтение {input_path.name}...")
        text = read_text_any(input_path)
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:                                   # один столбец — разделителя нет
            dialect = csv.excel
        # Потоковая запись: обычный режим держит в памяти всю таблицу (200 тыс. строк — +380 МБ)
        wb = Workbook(write_only=True)
        ws = wb.create_sheet()
        for n, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
            if n > EXCEL_MAX_ROWS:
                self._handle_error(f"В CSV больше {EXCEL_MAX_ROWS:,} строк — столько не помещается в лист Excel."
                                   .replace(",", " "))
                return False
            ws.append([_cell(value) for value in row])
        wb.save(output_path)
        self._update_progress(100)
        return True

    def _xlsx_to_csv(self, input_path: Path, output_path: Path) -> bool:
        from openpyxl import load_workbook

        from src.utils.helpers import get_unique_filename

        self._update_status("Сохранение в CSV...")
        wb = load_workbook(input_path, read_only=True, data_only=True)
        # CSV — это один лист. Первый — в выбранный файл, остальные — рядом: «отчёт - Февраль.csv».
        # Пустые листы («Лист2», «Лист3» из старых шаблонов) пропускаем
        sheets = [ws for ws in wb.worksheets
                  if any(v is not None for row in ws.iter_rows(values_only=True) for v in row)]
        delimiter = excel_csv_delimiter()
        comma = delimiter == ';'                                  # дробная часть — через запятую
        for n, ws in enumerate(sheets or wb.worksheets[:1]):
            title = re.sub(r'[<>:"/\\|?*]', '_', ws.title)        # недопустимое в именах файлов Windows
            target = output_path if n == 0 else \
                get_unique_filename(output_path.with_name(f"{output_path.stem} - {title}.csv"))
            # С BOM: без него русский Excel откроет UTF-8 кракозябрами
            with open(target, 'w', encoding='utf-8-sig', newline='') as f:
                csv.writer(f, delimiter=delimiter).writerows(
                    [_csv_value(v, comma) for v in row]
                    for row in ws.iter_rows(values_only=True))
        wb.close()
        self._update_progress(100)
        return True
