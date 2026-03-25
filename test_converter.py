"""
Тестовый скрипт для проверки конвертации
"""
from pathlib import Path
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet


def test_conversion(input_file, output_file):
    """Тестовая конвертация"""
    try:
        print(f"Чтение файла: {input_file}")
        doc = Document(input_file)

        print(f"Количество параграфов: {len(doc.paragraphs)}")

        # Печатаем первые 5 параграфов
        for i, para in enumerate(doc.paragraphs[:5]):
            print(f"Параграф {i}: {para.text[:100]}")

        print("Создание PDF...")
        pdf = SimpleDocTemplate(str(output_file), pagesize=A4)
        styles = getSampleStyleSheet()
        story = []

        for para in doc.paragraphs:
            if para.text.strip():
                story.append(Paragraph(para.text[:500], styles['Normal']))

        pdf.build(story)
        print(f"PDF создан: {output_file}")
        return True

    except Exception as e:
        print(f"Ошибка: {e}")
        import traceback
        traceback.print_exc()
        return False


if __name__ == "__main__":
    # Укажите путь к вашему файлу
    input_file = "/Users/ivt/Documents/ТЕХНИЧЕСКОЕ ЗАДАНИЕ_Тихонов.docx"
    output_file = "/Users/ivt/Downloads/test_output.pdf"

    test_conversion(Path(input_file), Path(output_file))